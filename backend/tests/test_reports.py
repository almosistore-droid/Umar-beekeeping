"""Stage 6: dashboard, yearly report and its PDF / Excel / CSV exports — on a hand-counted history."""
import csv
import io
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.services import telegram_bot

TODAY = datetime.now(timezone.utc).date()
Y = TODAY.year


def hive(client, headers, apiary_id, number, installed=None, family=True, **kw):
    body = {"apiary_id": apiary_id, "hive_number": number, **kw}
    if installed:
        body["installation_date"] = installed.isoformat()
    if family:
        body["family"] = {"strength": kw.pop("strength", "MEDIUM")}
    r = client.post("/api/hives", json=body, headers=headers)
    assert r.status_code == 201, r.text
    return r.json()


@pytest.fixture
def history(client, beekeeper):
    """
    Last year: A, B, D installed.            This year: C installed, D merged into A, B sold,
    a split of A creates E, queen of A replaced, honey A 10 kg @100 + C 5 kg, wax 2 kg @150,
    varroa found and cured in A, two treatments with Бипин.
    """
    ap = client.post("/api/apiaries", json={"name": "Пасека"}, headers=beekeeper).json()
    last_year = date(Y - 1, 3, 1)
    a = hive(client, beekeeper, ap["id"], "1", installed=last_year, strength="STRONG")
    b = hive(client, beekeeper, ap["id"], "2", installed=last_year)
    d = hive(client, beekeeper, ap["id"], "4", installed=last_year, strength="WEAK")
    c = hive(client, beekeeper, ap["id"], "3")

    q_old = client.post("/api/queens", json={"queen_code": "M-1", "hive_id": a["id"],
                                             "birth_date": (TODAY - timedelta(days=800)).isoformat()}, headers=beekeeper).json()
    q_new = client.post("/api/queens", json={"queen_code": "M-2", "birth_date": (TODAY - timedelta(days=100)).isoformat()},
                        headers=beekeeper).json()
    r = client.post(f"/api/hives/{a['id']}/queen", json={"queen_id": q_new["id"], "previous_queen_status": "REPLACED"},
                    headers=beekeeper)
    assert r.status_code == 200, r.text

    r = client.post("/api/merges", json={"source_hive_ids": [a["id"], d["id"]], "target_hive_id": a["id"],
                                         "keep_queen_id": q_new["id"]}, headers=beekeeper)
    assert r.status_code == 201, r.text
    sold = client.get(f"/api/hives/{b['id']}", headers=beekeeper).json()
    assert client.put(f"/api/hives/{b['id']}", json={"status": "SOLD", "version": sold["version"]},
                      headers=beekeeper).status_code == 200
    r = client.post(f"/api/hives/{a['id']}/split", json={"new_hive_number": "5"}, headers=beekeeper)
    assert r.status_code == 201, r.text
    e = r.json()["new_hive"]

    month = TODAY.replace(day=1)
    assert client.post(f"/api/hives/{a['id']}/honey", json={"weight_kg": 10, "price_per_kg": 100,
                                                            "harvest_date": month.isoformat()}, headers=beekeeper).status_code == 201
    assert client.post(f"/api/hives/{c['id']}/honey", json={"weight_kg": 5, "harvest_date": month.isoformat()},
                       headers=beekeeper).status_code == 201
    assert client.post("/api/products", json={"product_type": "WAX", "quantity": 2, "unit": "KG", "price_per_unit": 150},
                       headers=beekeeper).status_code == 201

    diseases = {x["code"]: x for x in client.get("/api/diseases", headers=beekeeper).json()}
    case = client.post(f"/api/hives/{a['id']}/diseases", json={"disease_id": diseases["VARROOSIS"]["id"],
                                                               "status": "CONFIRMED"}, headers=beekeeper).json()["case"]
    for _ in range(2):
        r = client.post(f"/api/hives/{a['id']}/treatments", json={"medication": "Бипин", "hive_disease_id": case["id"]},
                        headers=beekeeper)
        assert r.status_code == 201, r.text
    assert client.put(f"/api/disease-cases/{case['id']}", json={"status": "RESOLVED"}, headers=beekeeper).status_code == 200
    return {"a": a, "b": b, "c": c, "d": d, "e": e, "q_old": q_old, "q_new": q_new}


def report(client, headers, year=Y):
    r = client.get(f"/api/reports/yearly/{year}", headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


# ---------- yearly report ----------

def test_hives_section(client, beekeeper, history):
    h = report(client, beekeeper)["hives"]
    # start: A, B, D · new: C, E · gone: B sold, D merged · end: A, C, E
    assert h == {"at_start": 3, "at_end": 3, "new": 2, "lost": 0, "sold": 1, "merged": 1, "split": 1}
    prev = report(client, beekeeper, Y - 1)["hives"]
    assert prev["at_start"] == 0 and prev["new"] == 3 and prev["at_end"] == 3 and prev["sold"] == 0


def test_queens_section(client, beekeeper, history):
    q = report(client, beekeeper)["queens"]
    assert q["in_hives"] == 1 and q["new"] == 2 and q["replaced"] == 1
    assert Decimal(q["avg_age_months"]) == Decimal("3.0")  # M-2, born 100 days ago
    assert q["best_families"][0]["hive_number"] == "1" and q["best_families"][0]["queen_code"] == "M-2"


def test_honey_and_finance_sections(client, beekeeper, history):
    r = report(client, beekeeper)
    m = r["honey"]
    assert Decimal(m["total_kg"]) == 15 and m["producing_hives"] == 2 and Decimal(m["avg_per_hive_kg"]) == Decimal("7.5")
    assert m["best_hive"]["hive_number"] == "1" and m["worst_hive"]["hive_number"] == "3"
    assert Decimal(m["by_month"][TODAY.month - 1]) == 15 and sum(Decimal(x) for x in m["by_month"]) == 15
    assert Decimal(m["value"]) == 1000
    assert m["products"] == [{"product_type": "WAX", "unit": "KG", "quantity": "2.00", "value": "300.00"}]
    f = r["finance"]
    assert Decimal(f["production_value"]) == 1300 and f["visible"] is True and f["currency"] == "смн"
    assert Decimal(f["income"]) == 0 and Decimal(f["profit"]) == 0  # no sales or expenses recorded here


def test_diseases_section(client, beekeeper, history):
    d = report(client, beekeeper)["diseases"]
    assert d["cases"] == 1 and d["resolved"] == 1 and d["open"] == 0 and d["treatments"] == 2
    assert d["by_disease"] == [{"name": "Варроатоз", "cases": 1, "resolved": 1, "open": 0}]
    assert d["medications"] == [{"medication": "Бипин", "count": 2}]


def test_period_and_years(client, beekeeper, history):
    r = report(client, beekeeper)
    assert r["period_end"] == min(date(Y, 12, 31), TODAY).isoformat() and r["is_partial"] == (TODAY < date(Y, 12, 31))
    assert r["available_years"][:2] == [Y, Y - 1]


def test_empty_year_is_zeros(client, beekeeper):
    r = report(client, beekeeper, Y - 3)
    assert r["hives"]["at_end"] == 0 and Decimal(r["honey"]["total_kg"]) == 0 and r["honey"]["best_hive"] is None
    assert r["queens"]["avg_age_months"] is None and r["diseases"]["by_disease"] == []


@pytest.mark.parametrize("year", [Y + 1, 1999])
def test_bad_year(client, beekeeper, year):
    r = client.get(f"/api/reports/yearly/{year}", headers=beekeeper)
    assert r.status_code == 422 and "Год отчёта" in r.json()["detail"]


def test_report_needs_login_viewer_allowed(client, as_role):
    assert client.get(f"/api/reports/yearly/{Y}").status_code == 401
    assert client.get(f"/api/reports/yearly/{Y}", headers=as_role("VIEWER")).status_code == 200


# ---------- dashboard ----------

def test_dashboard(client, beekeeper, history):
    r = client.get("/api/dashboard", headers=beekeeper)
    assert r.status_code == 200, r.text
    d = r.json()
    t = d["tiles"]
    assert t["hives"] == 3 and t["queens_in_hives"] == 1 and t["queenless"] == 1  # E (split without a queen)
    assert Decimal(t["honey_year_kg"]) == 15 and t["best_hive"]["hive_number"] == "1"
    assert d["finance"]["currency"] == "смн" and len(d["finance"]["income_by_month"]) == 12
    assert [b["key"] for b in d["strength"]] == ["VERY_WEAK", "WEAK", "MEDIUM", "STRONG", "VERY_STRONG"]
    assert sum(b["count"] for b in d["strength"]) == 3
    assert {b["key"]: b["count"] for b in d["queen_ages"]}["lt1"] == 1
    assert d["diseases"] == [{"name": "Варроатоз", "cases": 1, "resolved": 1, "open": 0}]
    assert len(d["honey_by_month"]) == 12 and [h["hive_number"] for h in d["honey_by_hive"]] == ["1", "3"]


def test_dashboard_empty_and_year_check(client, beekeeper):
    d = client.get("/api/dashboard", headers=beekeeper).json()
    assert d["tiles"]["hives"] == 0 and d["honey_by_hive"] == [] and d["diseases"] == []
    assert client.get(f"/api/dashboard?year={Y + 1}", headers=beekeeper).status_code == 422


# ---------- exports ----------

def export(client, headers, fmt):
    r = client.get(f"/api/reports/yearly/{Y}/export", params={"format": fmt}, headers=headers)
    assert r.status_code == 200, r.text
    assert f'godovoy-otchet-{Y}.{fmt}' in r.headers["content-disposition"]
    return r


def test_export_pdf(client, beekeeper, history):
    r = export(client, beekeeper, "pdf")
    assert r.headers["content-type"] == "application/pdf" and r.content.startswith(b"%PDF")


def test_export_xlsx(client, beekeeper, history):
    r = export(client, beekeeper, "xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Сводка", "Мёд по ульям", "Мёд по месяцам", "Болезни"]
    summary = {row[0]: row[1] for row in wb["Сводка"].iter_rows(values_only=True) if row and row[0]}
    assert summary["Всего собрано, кг"] == "15" and summary["Проданы"] == "1"
    hives = list(wb["Мёд по ульям"].iter_rows(values_only=True))
    assert hives[0] == ("Место", "Улей", "Пасека", "кг", "Откачек") and hives[1][1:4] == ("№1", "Пасека", 10.0)
    assert list(wb["Болезни"].iter_rows(values_only=True))[1] == ("Варроатоз", 1, 1, 0)


def test_export_csv(client, beekeeper, history):
    r = export(client, beekeeper, "csv")
    assert r.content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][0].startswith(f"Годовой отчёт за {Y}")
    assert ["Мёд", "В среднем на улей, кг", "7,5"] in rows
    assert ["Болезни", "Препарат: Бипин", "2 раза"] in rows


def test_bad_format(client, beekeeper):
    assert client.get(f"/api/reports/yearly/{Y}/export", params={"format": "doc"}, headers=beekeeper).status_code == 422


def test_send_by_bot(client, beekeeper, history, monkeypatch):
    sent = {}

    def fake(chat_id, filename, content, caption=None, media_type="application/pdf"):
        sent.update(chat_id=chat_id, filename=filename, head=content[:2], caption=caption, media=media_type)
    monkeypatch.setattr(telegram_bot, "send_document", fake)
    r = client.post(f"/api/reports/yearly/{Y}/send", params={"format": "xlsx"}, headers=beekeeper)
    assert r.status_code == 200 and r.json()["sent"] is True
    assert sent["chat_id"] == 3000 and sent["filename"] == f"godovoy-otchet-{Y}.xlsx" and sent["head"] == b"PK"
    assert "Excel" in sent["caption"] and "spreadsheetml" in sent["media"]
