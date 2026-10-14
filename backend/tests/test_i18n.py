"""Тоҷикӣ on the server: catalog coverage, the template translator, X-Lang responses, saved language, files."""
import csv
import io
import json
import re
import uuid
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.i18n import translate
from app.i18n.extract import all_templates
from app.schemas.common import today as app_today

CATALOG = json.loads((Path(__file__).resolve().parents[1] / "app" / "i18n" / "tg.json").read_text(encoding="utf-8"))
TG = {"X-Lang": "tg"}


def test_every_server_text_has_a_tajik_translation():
    missing = sorted(t for t in all_templates() if t not in CATALOG)
    assert missing == [], "run: python -m app.i18n.extract — and add these to app/i18n/tg.json"


def test_translations_keep_placeholders_and_edge_spaces():
    ph = lambda s: sorted(re.findall(r"\{\d+\}", s))  # noqa: E731
    broken = [k for k, v in CATALOG.items() if ph(k) != ph(v) or k[:1].isspace() != v[:1].isspace()
              or k[-1:].isspace() != v[-1:].isspace()]
    assert broken == []


@pytest.mark.parametrize("ru, tg", [
    ("Улей не найден", "Қуттӣ ёфт нашуд"),
    ("Улей №27 уже есть на пасеке «Горная»", "Қуттии №27 аллакай дар занбӯрхонаи «Горная» ҳаст"),
    ("Проверьте введённые данные: Вес: должно быть больше 0", "Маълумоти воридшударо санҷед: Вазн: бояд аз 0 зиёд бошад"),
    ("Осмотр просрочен на 5 дней", "Муоина 5 рӯз дер монд"),
    ("≈ 2 года 3 месяца", "≈ 2 сол 3 моҳ"),
    ("Подозрение: Американский гнилец", "Гумон: Пӯсидагии амрикоӣ"),
    ("Матка не найдена и яиц нет — возможно, семья осталась без матки. Отметьте улей «Без матки» или подсадите матку.",
     "Модарзанбӯр ёфт нашуд ва тухм нест — шояд оила бе модарзанбӯр монда бошад. Қуттиро «Бе модарзанбӯр» қайд кунед ё модарзанбӯр гузоред."),
    ("Мой улей у реки", "Мой улей у реки"),  # user data stays as it is
    ("27", "27"),
])
def test_translate(ru, tg):
    assert translate(ru, "tg") == tg
    assert translate(ru, "ru") == ru


def test_errors_and_validation_in_tajik(client, beekeeper):
    r = client.get(f"/api/hives/{uuid.uuid4()}", headers={**beekeeper, **TG})
    assert r.status_code == 404 and r.json()["detail"] == "Қуттӣ ёфт нашуд"
    r = client.get(f"/api/hives/{uuid.uuid4()}", headers=beekeeper)
    assert r.json()["detail"] == "Улей не найден"
    ap = client.post("/api/apiaries", json={"name": "Ромит"}, headers=beekeeper).json()
    h = client.post("/api/hives", json={"apiary_id": ap["id"], "hive_number": "1", "family": {}}, headers=beekeeper).json()
    r = client.post(f"/api/hives/{h['id']}/honey", json={"weight_kg": 0}, headers={**beekeeper, **TG})
    assert r.status_code == 422
    assert r.json()["detail"] == "Маълумоти воридшударо санҷед: Вазн: бояд аз 0 зиёд бошад"
    assert r.json()["errors"][0]["message"] == "Вазн: бояд аз 0 зиёд бошад"


def test_data_is_not_translated_but_labels_are(client, beekeeper):
    ap = client.post("/api/apiaries", json={"name": "Ромит"}, headers=beekeeper).json()
    h = client.post("/api/hives", json={"apiary_id": ap["id"], "hive_number": "7"}, headers=beekeeper).json()
    check = client.get(f"/api/hives/{h['id']}/check", headers={**beekeeper, **TG}).json()
    assert check["apiary"]["name"] == "Ромит" and check["hive_number"] == "7"
    assert check["status"] == "EMPTY"  # codes stay codes
    assert check["issues"][0]["message"] == "Дар қуттӣ оилаи занбӯр нест"
    timeline = client.get(f"/api/hives/{h['id']}/timeline", headers={**beekeeper, **TG}).json()
    assert timeline  # translated without breaking the shape


def test_save_language(client, login):
    headers, body = login(4321)
    assert body["user"]["language"] is None
    r = client.put("/api/auth/me/language", json={"language": "tg"}, headers=headers)
    assert r.status_code == 200 and r.json()["language"] == "tg"
    _, again = login(4321)
    assert again["user"]["language"] == "tg"
    assert client.put("/api/auth/me/language", json={"language": "en"}, headers=headers).status_code == 422
    assert client.put("/api/auth/me/language", json={"language": "ru"}).status_code == 401


def test_report_files_in_tajik(client, beekeeper):
    y = app_today().year
    r = client.get(f"/api/reports/yearly/{y}/export", params={"format": "csv"}, headers={**beekeeper, **TG})
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][0].startswith(f"Ҳисоботи солона барои {y}")
    assert rows[2] == ["Бахш", "Нишондиҳанда", "Қимат"]
    assert ["Қуттиҳо", "Дар аввали сол", "0"] in rows
    r = client.get(f"/api/reports/yearly/{y}/export", params={"format": "xlsx"}, headers={**beekeeper, **TG})
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Хулоса", "Асал аз рӯи қуттиҳо", "Асал аз рӯи моҳҳо", "Бемориҳо"]
    r = client.get(f"/api/reports/yearly/{y}/export", params={"format": "pdf"}, headers={**beekeeper, **TG})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_json_with_x_lang_ru_or_missing_is_untouched(client, beekeeper):
    a = client.get("/api/dashboard", headers=beekeeper).json()
    b = client.get("/api/dashboard", headers={**beekeeper, "X-Lang": "ru"}).json()
    c = client.get("/api/dashboard", headers={**beekeeper, "X-Lang": "xx"}).json()
    assert a == b == c
    t = client.get("/api/dashboard", headers={**beekeeper, **TG}).json()
    assert [x["label"] for x in t["strength"]] == ["Хеле заиф", "Заиф", "Миёна", "Қавӣ", "Хеле қавӣ"]
